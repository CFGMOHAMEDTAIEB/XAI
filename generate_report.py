"""Generates rapport.docx - a full project report for XAI-Compress Platform."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ---------- Styles ----------
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# ---------- Title page ----------
title = doc.add_heading('XAI-Compress Platform', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph('Project Report')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(18)
sub.runs[0].font.color.rgb = RGBColor(0x40, 0x40, 0x40)
doc.add_paragraph('Monorepo MVP for a cross-platform lossless-compression ecosystem').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

# ---------- Table of Contents (Word field, updatable) ----------
doc.add_heading('Table of Contents', level=1)
paragraph = doc.add_paragraph()
run = paragraph.add_run()
fld_char = OxmlElement('w:fldChar')
fld_char.set(qn('w:fldCharType'), 'begin')
instr_text = OxmlElement('w:instrText')
instr_text.set(qn('xml:space'), 'preserve')
instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
fld_char2 = OxmlElement('w:fldChar')
fld_char2.set(qn('w:fldCharType'), 'separate')
fld_char3 = OxmlElement('w:t')
fld_char3.text = "Right-click and select 'Update Field' to generate the table of contents."
fld_char4 = OxmlElement('w:fldChar')
fld_char4.set(qn('w:fldCharType'), 'end')

r_element = run._r
r_element.append(fld_char)
r_element.append(instr_text)
r_element.append(fld_char2)
r_element.append(fld_char3)
r_element.append(fld_char4)
doc.add_page_break()

# ---------- 1. Executive Summary ----------
doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    "XAI-Compress Platform is a modular monorepo MVP implementing a cross-platform "
    "lossless-compression ecosystem. It combines a validated PyTorch + Rust compression "
    "engine with a working FastAPI backend, multiple frontend/backend service skeletons "
    "(Flutter, Angular, Next.js, .NET, Spring, Symfony, Node.js), and a Docker Compose "
    "infrastructure stack (PostgreSQL, Redis, MinIO, Keycloak, ClamAV, Nginx). "
    "The project is explicitly positioned as a working foundation/MVP rather than a "
    "production-ready, bank-grade platform: some modules are fully executable, others "
    "are documented extension points reserved for future sprints."
)
doc.add_paragraph(
    "The core value proposition is a secure, authenticated file-compression and sharing "
    "workflow: users register and enroll in TOTP-based multi-factor authentication, "
    "compress files locally through the AI-driven engine, optionally sync metadata to "
    "the cloud, and share files through time-limited, hashed, one-time share codes."
)

# ---------- 2. Project Structure ----------
doc.add_heading('2. Project Structure', level=1)
doc.add_paragraph('Top-level directories and their purpose:')
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Directory'
hdr[1].text = 'Purpose'
rows = [
    ('apps/', 'Frontend applications: desktop, mobile, web, and admin clients'),
    ('services/', 'Backend services: FastAPI API, Node realtime, Spring, Symfony'),
    ('engines/', 'AI compression engine (PyTorch + Rust) - "XAI-Compress" core'),
    ('packages/', 'Shared contracts (Protobuf) used across services'),
    ('infrastructure/', 'Docker Compose, Keycloak realm, Nginx, Kubernetes, Vault, monitoring'),
    ('analytics/', 'Benchmark, visualization, and reporting suite'),
    ('labs/', 'Experimental PoCs: C++ SIMD, NLP/CV/RAG, TensorFlow comparison'),
    ('tools/', 'Developer tooling (Qt container inspector)'),
    ('scripts/', 'Prerequisite checks and structure validation scripts'),
    ('docs/', 'Architecture, roadmap, and security documentation'),
]
for name, purpose in rows:
    row = table.add_row().cells
    row[0].text = name
    row[1].text = purpose

# ---------- 3. Technology Stack ----------
doc.add_heading('3. Technology Stack', level=1)

doc.add_heading('3.1 Backend Services', level=2)
t2 = doc.add_table(rows=1, cols=3)
t2.style = 'Light Grid Accent 1'
h = t2.rows[0].cells
h[0].text = 'Service'
h[1].text = 'Stack'
h[2].text = 'Status'
backend_rows = [
    ('api_fastapi', 'Python 3.12, FastAPI, SQLAlchemy, PostgreSQL/SQLite, Redis, JWT, TOTP (pyotp)', 'Executable / working MVP'),
    ('realtime_node', 'Node.js', 'Extension point (skeleton)'),
    ('enterprise_spring', 'Java / Spring', 'Extension point (skeleton)'),
    ('support_symfony', 'PHP / Symfony', 'Extension point (skeleton)'),
]
for a, b, c in backend_rows:
    r = t2.add_row().cells
    r[0].text, r[1].text, r[2].text = a, b, c

doc.add_heading('3.2 Frontend Applications', level=2)
t3 = doc.add_table(rows=1, cols=3)
t3.style = 'Light Grid Accent 1'
h = t3.rows[0].cells
h[0].text = 'App'
h[1].text = 'Stack'
h[2].text = 'Status'
frontend_rows = [
    ('desktop_flutter', 'Flutter / Dart', 'Source skeleton'),
    ('mobile_authenticator_flutter', 'Flutter / Dart (TOTP, 30s codes)', 'Source skeleton'),
    ('web_angular', 'Angular / TypeScript', 'Source skeleton'),
    ('public_nextjs', 'Next.js / React', 'Extension point'),
    ('admin_dotnet', '.NET', 'Extension point'),
]
for a, b, c in frontend_rows:
    r = t3.add_row().cells
    r[0].text, r[1].text, r[2].text = a, b, c

doc.add_heading('3.3 Core Engine', level=2)
doc.add_paragraph(
    "engines/ai_compression (XAI-Compress): a previously validated lossless compression "
    "engine combining PyTorch (model-driven compression) and a Rust extension for "
    "performance-critical routines. Produces .xaic container files, verified with SHA-256 "
    "integrity checks and corrupted-container rejection."
)

doc.add_heading('3.4 Infrastructure (Docker Compose)', level=2)
t4 = doc.add_table(rows=1, cols=2)
t4.style = 'Light Grid Accent 1'
h = t4.rows[0].cells
h[0].text = 'Component'
h[1].text = 'Role'
infra_rows = [
    ('PostgreSQL 17', 'Primary data store: users, file metadata, shares, audits'),
    ('Redis 7 (alpine)', 'Rate limiting, short-lived state, token revocation, job progress'),
    ('MinIO', 'S3-compatible encrypted object storage (not yet wired to API)'),
    ('Keycloak 26.3', 'Target identity provider / OIDC (dev mode with realm import)'),
    ('ClamAV', 'Malware scanning for uploaded files'),
    ('Nginx', 'Reverse proxy (infrastructure placeholder)'),
    ('Grafana / ELK', 'Monitoring and logging (placeholders)'),
    ('Kubernetes / Vault', 'Future orchestration and secrets management (placeholders)'),
]
for a, b in infra_rows:
    r = t4.add_row().cells
    r[0].text, r[1].text = a, b

# ---------- 4. Architecture & Data Flows ----------
doc.add_heading('4. Architecture & Data Flows', level=1)
doc.add_paragraph(
    "Each module runs independently and integrates through explicit contracts "
    "(Protobuf for internal services, REST for the API). Three vertical flows define "
    "the MVP behavior:"
)
doc.add_heading('4.1 Authentication Flow', level=2)
doc.add_paragraph('Register → enroll TOTP → confirm code → login with password + TOTP → receive token.')
doc.add_heading('4.2 Sharing Flow', level=2)
doc.add_paragraph(
    'Record compressed file metadata → create recipient-bound one-time share code → '
    'deliver code externally → recipient authenticates → redeems code → sees file '
    'metadata → storage integration supplies encrypted download.'
)
doc.add_heading('4.3 Local Desktop Compression Flow', level=2)
doc.add_paragraph(
    'Select file → Flutter launches local Python CLI → engine creates .xaic → Flutter '
    'presents metrics → API stores history metadata if user chooses cloud sync.'
)

# ---------- 5. Security Posture ----------
doc.add_heading('5. Security Posture', level=1)
doc.add_heading('5.1 Implemented in the MVP', level=2)
for item in [
    'Password hashing', 'Short-lived JWTs', 'TOTP enrollment and verification',
    'Cryptographically random, server-side hashed share codes',
    'Share code expiry and download limits', 'Recipient binding', 'Audit events',
    'Engine SHA-256 integrity checks and corrupted-container rejection',
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('5.2 Required Before Production', level=2)
for item in [
    'Exclusive use of Keycloak for credentials and MFA',
    'Argon2id password policy and breached-password screening',
    'TLS everywhere and certificate pinning where appropriate',
    'Redis-backed rate limits and distributed locks',
    'AES-256-GCM envelope encryption with Vault-managed keys',
    'MinIO private buckets with short-lived signed URLs',
    'ClamAV/YARA scan orchestration, quarantine and sandbox policy',
    'Refresh-token rotation and device revocation',
    'CSRF protection for browser flows',
    'Secure email provider with anti-enumeration responses',
    'Privacy, retention, deletion, and incident-response policies',
    'External penetration test',
]:
    doc.add_paragraph(item, style='List Bullet')

# ---------- 6. Roadmap ----------
doc.add_heading('6. Roadmap', level=1)
doc.add_heading('6.1 Delivered Foundation', level=2)
for item in [
    'Validated lossless engine with Rust extension', 'Benchmark suite',
    'FastAPI account/TOTP/share/history MVP',
    'Flutter desktop and authenticator source skeletons',
    'Angular portal source skeleton', 'Docker infrastructure', 'Protobuf contracts',
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('6.2 Next Sprint', level=2)
for item in [
    'Generate Flutter platform folders and implement API clients',
    'Replace manual mobile secret entry with QR scanning',
    'Connect all clients to Keycloak OIDC',
    'Integrate MinIO storage and AES-GCM encryption',
    'Add email delivery for share codes',
    'Add Redis rate limits and background jobs',
]:
    doc.add_paragraph(item, style='List Number')

doc.add_heading('6.3 Later Modules', level=2)
for item in [
    'ClamAV/YARA orchestration', 'Angular admin portal', 'Node WebSocket gateway',
    'Kafka events', 'ONNX deployment', 'C++ SIMD benchmark laboratory',
    'Qt container inspector', '.NET security console', 'Spring/Symfony enterprise PoCs',
    'Vault, ELK, Grafana, SonarQube, Trivy, GitLab CI, Kubernetes',
    'NLP/CV/RAG optional PoCs with strict privacy boundaries',
]:
    doc.add_paragraph(item, style='List Bullet')

# ---------- 7. Getting Started ----------
doc.add_heading('7. Getting Started', level=1)
doc.add_heading('7.1 Backend (FastAPI)', level=2)
code = doc.add_paragraph()
code.add_run(
    "cd services/api_fastapi\n"
    "python -m venv .venv\n"
    "pip install -r requirements.txt\n"
    "copy .env.example .env\n"
    "uvicorn app.main:app --reload --port 8000"
).font.name = 'Consolas'
doc.add_paragraph('Then open http://localhost:8000/docs')

doc.add_heading('7.2 Docker Infrastructure', level=2)
code2 = doc.add_paragraph()
code2.add_run('docker compose up -d postgres redis minio keycloak').font.name = 'Consolas'

doc.add_heading('7.3 Default Local Keycloak Admin', level=2)
doc.add_paragraph('username: admin  |  password: change-me-now  (must be changed before any deployment)')

# ---------- 8. Conclusion ----------
doc.add_heading('8. Conclusion', level=1)
doc.add_paragraph(
    "XAI-Compress Platform delivers a coherent, working MVP foundation: a validated "
    "compression engine, a functional authenticated REST API, and a containerized "
    "supporting infrastructure. The remaining modules (mobile/web clients, enterprise "
    "connectors, monitoring, orchestration) are clearly scoped as extension points with "
    "an explicit roadmap, avoiding overstated production readiness while providing a "
    "credible path from MVP to a hardened, deployable platform."
)

doc.save(r'C:\Users\ss\Desktop\XAI\XAI\rapport.docx')
print('done')
